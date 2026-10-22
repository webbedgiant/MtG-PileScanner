"""Spreadsheet building/export."""
import csv, io, os
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from . import db

COLUMNS = ["Qty", "Foil", "Name", "Set", "Set Name", "Collector #", "Rarity", "Mana Cost", "Mana Value",
           "Type", "Colors", "Power", "Toughness", "Price (USD)", "Total (USD)", "Released", "Scryfall", "Added"]


def export_path():
    p = db.meta_get("setting:export_path")
    return Path(p) if p else db.DATA_DIR / "collection.xlsx"


def fetch():
    """Returns (keys, rows) where rows align with COLUMNS."""
    with db.LOCK:
        recs = db.get().execute(
            """SELECT c.card_id, c.foil, c.qty, c.added_at, k.* FROM collection c
               JOIN cards k ON k.id = c.card_id
               ORDER BY k.name COLLATE NOCASE, k.set_code, c.foil""").fetchall()
    keys, rows = [], []
    for r in recs:
        price = r["price_usd_foil"] if r["foil"] and r["price_usd_foil"] is not None else r["price_usd"]
        keys.append({"card_id": r["card_id"], "foil": r["foil"], "image_url": r["image_url"]})
        rows.append([r["qty"], "Foil" if r["foil"] else "", r["name"], (r["set_code"] or "").upper(),
                     r["set_name"], r["collector_number"], r["rarity"], r["mana_cost"], r["cmc"],
                     r["type_line"], "/".join(r["colors"] or ""), r["power"], r["toughness"],
                     price, round(price * r["qty"], 2) if price is not None else None,
                     r["released_at"], r["scryfall_uri"], (r["added_at"] or "")[:10]])
    return keys, rows


def workbook():
    _, rows = fetch()
    wb = Workbook()
    ws = wb.active
    ws.title = "Collection"
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
        c.alignment = Alignment(horizontal="center")
    widths = [6, 7, 34, 7, 28, 11, 10, 14, 11, 34, 9, 7, 10, 12, 12, 12, 40, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 14).number_format = ws.cell(r, 15).number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    s = wb.create_sheet("Summary")
    s.append(["Total cards", "=SUM(Collection!A:A)"])
    s.append(["Unique printings", "=COUNTA(Collection!C:C)-1"])
    s.append(["Estimated value (USD)", "=SUM(Collection!O:O)"])
    s["B3"].number_format = '"$"#,##0.00'
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 16
    return wb


def write_xlsx(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    workbook().save(tmp)
    os.replace(tmp, path)  # raises PermissionError if the file is open in Excel


def xlsx_bytes():
    buf = io.BytesIO()
    workbook().save(buf)
    return buf.getvalue()


def csv_bytes():
    _, rows = fetch()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel reads accents correctly


def backup(keep=14):
    """Daily CSV snapshot of the collection in <data folder>/backups (newest `keep` are retained)."""
    _, rows = fetch()
    if not rows:
        return
    d = db.DATA_DIR / "backups"
    d.mkdir(exist_ok=True)
    (d / f"collection-{date.today().isoformat()}.csv").write_bytes(csv_bytes())
    for old in sorted(d.glob("collection-*.csv"))[:-keep]:
        old.unlink()
